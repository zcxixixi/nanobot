#!/usr/bin/env python3
"""View holdings from Asset-Management workbook."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REQUIRED_COLUMNS = [
    "timestamp",
    "account",
    "symbol",
    "name",
    "quantity",
    "price_usd",
    "market_value_usd",
]


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[5]


def _default_workbook() -> Path:
    primary = _repo_root() / "Asset-Management" / "assets.xlsx"
    if primary.exists():
        return primary
    return Path.cwd() / "assets.xlsx"


def _resolve_workbook(path_arg: str | None) -> Path:
    candidate = Path(path_arg).expanduser() if path_arg else _default_workbook()
    resolved = candidate.resolve()
    if not resolved.exists():
        raise FileNotFoundError(f"Workbook not found: {resolved}")
    return resolved


def _header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=col).value
        if value is None:
            continue
        key = str(value).strip().lower()
        if key:
            mapping[key] = col
    return mapping


def _to_float(value: Any) -> float:
    if value is None:
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _read_holdings(workbook_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(workbook_path, data_only=True)
    if "Holdings" not in wb.sheetnames:
        raise ValueError("Sheet 'Holdings' not found")

    ws = wb["Holdings"]
    cols = _header_map(ws)
    missing = [name for name in REQUIRED_COLUMNS if name not in cols]
    if missing:
        raise ValueError(f"Holdings missing required columns: {missing}")

    rows: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        symbol = ws.cell(row=row_idx, column=cols["symbol"]).value
        if symbol is None or not str(symbol).strip():
            continue

        entry = {
            "timestamp": str(ws.cell(row=row_idx, column=cols["timestamp"]).value or ""),
            "account": str(ws.cell(row=row_idx, column=cols["account"]).value or ""),
            "symbol": str(symbol).strip().upper(),
            "name": str(ws.cell(row=row_idx, column=cols["name"]).value or ""),
            "quantity": _to_float(ws.cell(row=row_idx, column=cols["quantity"]).value),
            "price_usd": _to_float(ws.cell(row=row_idx, column=cols["price_usd"]).value),
            "market_value_usd": _to_float(ws.cell(row=row_idx, column=cols["market_value_usd"]).value),
        }
        rows.append(entry)

    return rows


def _render_table(rows: list[dict[str, Any]]) -> str:
    headers = ["symbol", "name", "account", "quantity", "price_usd", "market_value_usd"]
    data_rows = []
    for row in rows:
        data_rows.append(
            [
                row["symbol"],
                row["name"],
                row["account"],
                f"{row['quantity']:.6f}".rstrip("0").rstrip("."),
                f"{row['price_usd']:.4f}",
                f"{row['market_value_usd']:.2f}",
            ]
        )

    widths = [len(h) for h in headers]
    for item in data_rows:
        for i, value in enumerate(item):
            widths[i] = max(widths[i], len(value))

    line = " | ".join(h.ljust(widths[i]) for i, h in enumerate(headers))
    sep = "-+-".join("-" * widths[i] for i in range(len(headers)))

    out = [line, sep]
    for item in data_rows:
        out.append(" | ".join(value.ljust(widths[i]) for i, value in enumerate(item)))
    return "\n".join(out)


def main() -> int:
    parser = argparse.ArgumentParser(description="View holdings from assets.xlsx")
    parser.add_argument("--workbook", type=str, default=None, help="Path to workbook")
    parser.add_argument("--json", action="store_true", help="Print JSON output")
    args = parser.parse_args()

    try:
        workbook_path = _resolve_workbook(args.workbook)
        rows = _read_holdings(workbook_path)
    except Exception as exc:
        print(f"view_holdings failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1

    total = sum(float(item["market_value_usd"]) for item in rows)

    if args.json:
        payload = {
            "workbook": str(workbook_path),
            "count": len(rows),
            "total_market_value_usd": round(total, 2),
            "holdings": rows,
        }
        print(json.dumps(payload, ensure_ascii=False, indent=2))
    else:
        print(f"Workbook: {workbook_path}")
        print(f"Holdings count: {len(rows)}")
        print(f"Total market value (USD): {total:.2f}")
        if rows:
            print()
            print(_render_table(rows))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
