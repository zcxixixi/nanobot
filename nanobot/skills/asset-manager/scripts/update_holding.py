#!/usr/bin/env python3
"""Upsert a holding row in Asset-Management workbook by symbol."""

from __future__ import annotations

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REQUIRED_COLUMNS = [
    "timestamp",
    "account",
    "symbol",
    "name",
    "quantity",
    "price_usd",
    "market_value_usd",
]


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[5]


def _default_workbook() -> Path:
    primary = _repo_root() / "Asset-Management" / "assets.xlsx"
    if primary.exists():
        return primary
    return Path.cwd() / "assets.xlsx"


def _resolve_workbook(path_arg: str | None) -> Path:
    candidate = Path(path_arg).expanduser() if path_arg else _default_workbook()
    resolved = candidate.resolve()
    if not resolved.exists():
        raise FileNotFoundError(f"Workbook not found: {resolved}")
    return resolved


def _header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=col).value
        if value is None:
            continue
        key = str(value).strip().lower()
        if key:
            mapping[key] = col
    return mapping


def _to_float(value: Any) -> float:
    if value is None:
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _normalize_symbol(symbol: str) -> str:
    return str(symbol or "").strip().upper()


def _find_symbol_rows(ws, symbol_col: int, symbol: str) -> list[int]:
    rows: list[int] = []
    for row_idx in range(2, ws.max_row + 1):
        raw = ws.cell(row=row_idx, column=symbol_col).value
        if _normalize_symbol(raw) == symbol:
            rows.append(row_idx)
    return rows


def _infer_default_account(ws, account_col: int) -> str:
    for row_idx in range(2, ws.max_row + 1):
        raw = ws.cell(row=row_idx, column=account_col).value
        if raw is not None and str(raw).strip():
            return str(raw).strip()
    return "uSmart"


def _create_backup(path: Path) -> Path:
    backup_dir = path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_path = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
    shutil.copy2(path, backup_path)

    existing = sorted(backup_dir.glob(f"{path.stem}_*.xlsx"), key=lambda p: p.stat().st_mtime)
    for old in existing[:-10]:
        old.unlink(missing_ok=True)

    return backup_path


def main() -> int:
    parser = argparse.ArgumentParser(description="Add or update a holding by symbol")
    parser.add_argument("--symbol", required=True, help="Ticker symbol, e.g. AAPL.US")
    parser.add_argument("--quantity", required=True, type=float, help="Absolute quantity to set")
    parser.add_argument("--name", default=None, help="Optional asset display name")
    parser.add_argument("--account", default=None, help="Optional account name")
    parser.add_argument("--price-usd", default=None, type=float, help="Optional latest unit price")
    parser.add_argument("--workbook", default=None, help="Path to workbook")
    args = parser.parse_args()

    symbol = _normalize_symbol(args.symbol)
    if not symbol:
        print("update_holding failed: symbol must be non-empty", file=sys.stderr)
        return 1

    workbook_path = _resolve_workbook(args.workbook)

    try:
        wb = load_workbook(workbook_path)
        if "Holdings" not in wb.sheetnames:
            raise ValueError("Sheet 'Holdings' not found")

        ws = wb["Holdings"]
        cols = _header_map(ws)
        missing = [name for name in REQUIRED_COLUMNS if name not in cols]
        if missing:
            raise ValueError(f"Holdings missing required columns: {missing}")

        symbol_rows = _find_symbol_rows(ws, cols["symbol"], symbol)
        operation = "updated" if symbol_rows else "created"
        target_row = symbol_rows[0] if symbol_rows else ws.max_row + 1

        existing_price = _to_float(ws.cell(row=target_row, column=cols["price_usd"]).value)
        price = float(args.price_usd) if args.price_usd is not None else existing_price

        existing_name = str(ws.cell(row=target_row, column=cols["name"]).value or "").strip()
        existing_account = str(ws.cell(row=target_row, column=cols["account"]).value or "").strip()

        name = args.name if args.name is not None else (existing_name or symbol)
        account = args.account if args.account is not None else (existing_account or _infer_default_account(ws, cols["account"]))
        quantity = float(args.quantity)
        market_value = quantity * price

        ws.cell(row=target_row, column=cols["timestamp"], value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        ws.cell(row=target_row, column=cols["account"], value=account)
        ws.cell(row=target_row, column=cols["symbol"], value=symbol)
        ws.cell(row=target_row, column=cols["name"], value=name)
        ws.cell(row=target_row, column=cols["quantity"], value=quantity)
        ws.cell(row=target_row, column=cols["price_usd"], value=price)
        ws.cell(row=target_row, column=cols["market_value_usd"], value=market_value)

        backup_path = _create_backup(workbook_path)
        wb.save(workbook_path)

    except Exception as exc:
        print(f"update_holding failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1

    print(
        f"{operation} symbol={symbol} quantity={quantity:.6f} price_usd={price:.4f} "
        f"market_value_usd={market_value:.2f} workbook={workbook_path} backup={backup_path}"
    )
    if len(symbol_rows) > 1:
        print(f"warning: duplicate symbol rows detected for {symbol}, updated first row only", file=sys.stderr)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
