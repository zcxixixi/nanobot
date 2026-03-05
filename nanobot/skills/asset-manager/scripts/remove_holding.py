#!/usr/bin/env python3
"""Remove holding rows from Asset-Management workbook by symbol."""

from __future__ import annotations

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[5]


def _default_workbook() -> Path:
    primary = _repo_root() / "Asset-Management" / "assets.xlsx"
    if primary.exists():
        return primary
    return Path.cwd() / "assets.xlsx"


def _resolve_workbook(path_arg: str | None) -> Path:
    candidate = Path(path_arg).expanduser() if path_arg else _default_workbook()
    resolved = candidate.resolve()
    if not resolved.exists():
        raise FileNotFoundError(f"Workbook not found: {resolved}")
    return resolved


def _normalize_symbol(symbol: str) -> str:
    return str(symbol or "").strip().upper()


def _header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=col).value
        if value is None:
            continue
        key = str(value).strip().lower()
        if key:
            mapping[key] = col
    return mapping


def _create_backup(path: Path) -> Path:
    backup_dir = path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_path = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
    shutil.copy2(path, backup_path)

    existing = sorted(backup_dir.glob(f"{path.stem}_*.xlsx"), key=lambda p: p.stat().st_mtime)
    for old in existing[:-10]:
        old.unlink(missing_ok=True)

    return backup_path


def main() -> int:
    parser = argparse.ArgumentParser(description="Remove holding rows by symbol")
    parser.add_argument("--symbol", required=True, help="Ticker symbol, e.g. AAPL.US")
    parser.add_argument("--workbook", default=None, help="Path to workbook")
    args = parser.parse_args()

    symbol = _normalize_symbol(args.symbol)
    if not symbol:
        print("remove_holding failed: symbol must be non-empty", file=sys.stderr)
        return 1

    try:
        workbook_path = _resolve_workbook(args.workbook)
        wb = load_workbook(workbook_path)
        if "Holdings" not in wb.sheetnames:
            raise ValueError("Sheet 'Holdings' not found")

        ws = wb["Holdings"]
        cols = _header_map(ws)
        if "symbol" not in cols:
            raise ValueError("Holdings missing required column: symbol")

        symbol_col = cols["symbol"]
        target_rows: list[int] = []
        for row_idx in range(2, ws.max_row + 1):
            raw = ws.cell(row=row_idx, column=symbol_col).value
            if _normalize_symbol(raw) == symbol:
                target_rows.append(row_idx)

        if not target_rows:
            print(f"remove_holding: symbol not found: {symbol}", file=sys.stderr)
            return 1

        backup_path = _create_backup(workbook_path)
        for row_idx in sorted(target_rows, reverse=True):
            ws.delete_rows(row_idx, 1)

        wb.save(workbook_path)

    except Exception as exc:
        print(f"remove_holding failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1

    print(
        f"removed symbol={symbol} rows_deleted={len(target_rows)} "
        f"workbook={workbook_path} backup={backup_path}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
